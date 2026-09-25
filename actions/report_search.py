""""What did my report say about deep sleep?" -- saved reports, searched by meaning.

Research and market reports are saved as Word files in the JARVIS folder
(and moved into its Reports folder by the tidy-up). This finds the passage
that answers a question and reads it out, saying which report it came
from and when:

    what did my report say about deep sleep
    what does my ESP32 report say about batteries   (one report only)
    search my reports for energy use
    which reports mention insurance

Free and local: the passage is found by meaning with the same model as
memory and notes, each passage encoded once into the vector store, and it
is read out as written -- no model is asked to summarise it.

Reports are split by their own sections. A passage never runs across a
heading, and carries its heading when it is encoded, so "memory cards"
finds the paragraph under "Storage and memory cards" rather than whatever
followed it. The Sources list at the end is left out: it is URLs.
"""

import os
import re
import threading
import time
from datetime import date, datetime

from actions import files


# ---- recognising the question ---------------------------------------------

_ABOUT = r"(?:about|on|regarding|concerning|for|with\s+regard\s+to)"

# "my report", "the esp32 report", "my bitcoin vs ethereum research report"
_WHICH_REPORT = (
    r"(?:my|the|that|our)\s+(?:(?P<subject>.+?)\s+)?(?:research\s+|market\s+)?"
    r"(?:reports?|research)"
)

_SAY = r"(?:say|said|says|find|found|have|has|tell\s+me|mention|mentions|cover|covers)"

_QUESTIONS = (
    ("say", re.compile(
        rf"^what\s+(?:did|does|do|has|have)\s+{_WHICH_REPORT}\s+{_SAY}\s+{_ABOUT}\s+(?P<topic>.+)$")),
    ("say", re.compile(
        rf"^what\s+(?:is|s)\s+in\s+{_WHICH_REPORT}\s+{_ABOUT}\s+(?P<topic>.+)$")),
    ("say", re.compile(
        rf"^(?:search|check|look\s+through|look\s+in|go\s+through)\s+{_WHICH_REPORT}\s+"
        rf"(?:for|about)\s+(?P<topic>.+)$")),
    ("say", re.compile(
        rf"^(?:find|look\s+up)\s+(?P<topic>.+?)\s+in\s+{_WHICH_REPORT}$")),
    ("which", re.compile(
        r"^(?:which|what)\s+(?:of\s+my\s+)?reports?\s+"
        r"(?:mention|mentions|mentioned|talk\s+about|talks\s+about|cover|covers|covered|discuss|discusses|"
        r"say\s+anything\s+about|are\s+about|is\s+about)\s+(?P<topic>.+)$")),
    ("which", re.compile(
        r"^(?:do|does|did)\s+(?:any\s+of\s+)?(?:my|the)\s+reports?\s+"
        r"(?:mention|talk\s+about|cover|discuss|say\s+anything\s+about)\s+(?P<topic>.+)$")),
)

_EMPTY = frozenset({"it", "that", "this", "them", "anything", "something", "everything"})


def _normalise(said):
    text = str(said or "").casefold().replace("'", "").replace("\u2019", "")
    text = " ".join(re.sub(r"[^\w\s]", " ", text).split())

    previous = None

    while previous != text:
        previous = text
        text = re.sub(r"^(?:jarvis|hey|please|ok|okay)\s+", "", text)

    return text


def question(said):
    """(mode, topic, report subject or None) for a question about reports, or None."""
    text = _normalise(said)

    for mode, pattern in _QUESTIONS:
        match = pattern.match(text)

        if not match:
            continue

        topic = match.group("topic").strip()
        groups = match.groupdict()
        subject = (groups.get("subject") or "").strip() or None

        if subject in ("research", "market", "latest", "last", "recent", "saved"):
            subject = None

        if not topic or topic in _EMPTY:
            return None

        return mode, topic, subject

    return None


# ---- the reports -------------------------------------------------------------

# Reports are looked for in the JARVIS folder and one level of subfolders:
# Reports, and any folder a report was asked to be saved into.
_STAMP = re.compile(r"\s*(?:research\s+)?report\s+(\d{4}-\d{2}-\d{2})(?:-\d{6})?(?:\s*\(\d+\))?$", re.IGNORECASE)

_cache = {}
_cache_lock = threading.Lock()


def report_files():
    """Word files with "report" in their name, in the JARVIS folder and its subfolders."""
    base = files.root()

    if not base:
        return []

    folders = [base]

    try:
        folders += [
            entry.path for entry in os.scandir(base)
            if entry.is_dir() and not entry.name.startswith(".")
        ]
    except OSError:
        pass

    found = []

    for folder in folders:
        try:
            names = os.listdir(folder)
        except OSError:
            continue

        found += [
            os.path.join(folder, name) for name in names
            if name.casefold().endswith(".docx")
            and "report" in name.casefold()
            and not name.startswith("~$")
        ]

    return sorted(found)


def describe_report(path):
    """(what it is about, the day it was written) from its file name."""
    stem = os.path.splitext(os.path.basename(path))[0]
    match = _STAMP.search(stem)

    if match:
        label = stem[:match.start()].strip()

        try:
            day = datetime.strptime(match.group(1), "%Y-%m-%d").date()
        except ValueError:
            day = None
    else:
        label = re.sub(r"\s*report\s*$", "", stem, flags=re.IGNORECASE).strip()
        day = None

    if day is None:
        try:
            day = datetime.fromtimestamp(os.path.getmtime(path)).date()
        except OSError:
            day = None

    return label or stem, day


_HEADING_MARK = re.compile(r"^\s*#{1,6}\s*")
_BOLD_LINE = re.compile(r"^\s*\*\*(.+?)\*\*\s*:?\s*$")
_BULLET = re.compile(r"^\s*(?:[-*\u2022]|\d+[.)])\s+")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}")

# Passages are paragraphs, joined while short, never across a heading.
_PASSAGE_CHARS = 420


def _clean(line):
    line = line.replace("**", "").replace("__", "").replace("`", "")

    if "|" in line:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        line = ", ".join(cell for cell in cells if cell)

    return " ".join(line.split())


def _is_heading(raw, style):
    if style and style.casefold().startswith(("heading", "title")):
        return True

    if _HEADING_MARK.match(raw) or _BOLD_LINE.match(raw):
        return True

    # A table row is short and capitalised too, but never a heading.
    if "|" in raw:
        return False

    text = _clean(raw)

    return (
        0 < len(text) <= 70
        and not _BULLET.match(raw)
        and not text.endswith((".", ":", ";", ",", "?", "!"))
        and len(text.split()) <= 9
        and text[0].isupper()
    )


def sections(lines):
    """[(heading, passage)] from a report's lines, stopping at its Sources.

    [lines] are (text, style name or None) pairs, one per Word paragraph.
    """
    heading = ""
    found = []
    current = []
    columns = None

    def flush():
        if current:
            found.append((heading, " ".join(current)))
            current.clear()

    for raw, style in lines:
        raw = str(raw or "")

        if not raw.strip() or _TABLE_RULE.match(raw):
            continue

        if "|" in raw:
            cells = [_clean(cell) for cell in raw.strip().strip("|").split("|")]

            if columns is None:
                # The first row names the columns; later rows are read with
                # them, which is how a table makes sense aloud.
                columns = cells
                continue

            pairs = [
                f"{name}: {value}" if name else value
                for name, value in zip(columns + [""] * len(cells), cells)
                if value
            ]
            text = ", ".join(pairs) + "."

            if current and len(" ".join(current)) + 1 + len(text) > _PASSAGE_CHARS:
                flush()

            current.append(text)
            continue

        columns = None

        if _is_heading(raw, style):
            flush()
            heading = _clean(_HEADING_MARK.sub("", raw)).strip(" :")

            if heading.casefold() in ("sources", "references", "bibliography"):
                return found

            continue

        text = _clean(_BULLET.sub("", raw))

        if not text:
            continue

        if current and len(" ".join(current)) + 1 + len(text) > _PASSAGE_CHARS:
            flush()

        current.append(text)

    flush()

    return found


def _read(path):
    """The report's (heading, passage) pairs, cached until the file changes."""
    try:
        status = os.stat(path)
    except OSError:
        return []

    key = (status.st_mtime_ns, status.st_size)

    with _cache_lock:
        cached = _cache.get(path)

        if cached and cached[0] == key:
            return cached[1]

    try:
        from docx import Document

        document = Document(path)
        lines = [(paragraph.text, getattr(paragraph.style, "name", None)) for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                lines.append((" | ".join(cell.text for cell in row.cells), None))
    except Exception as error:
        print(f"[JARVIS] could not read the report {os.path.basename(path)}: {error}")
        return []

    found = sections(lines)

    with _cache_lock:
        _cache[path] = (key, found)

    return found


def passages(paths=None):
    """Every searchable passage, as (path, heading, passage)."""
    found = []

    for path in paths if paths is not None else report_files():
        for heading, passage in _read(path):
            found.append((path, heading, passage))

    return found


def _encoded(heading, passage):
    """What is encoded for a passage: its heading gives it context."""
    return f"{heading}: {passage}" if heading else passage


# ---- matching --------------------------------------------------------------

# Measured on research-report passages with the local model: passages about
# the topic scored 0.34 or more, unrelated topics 0.20 or less. A shared
# word adds 0.25, for the short topics the model underrates.
_MINIMUM = 0.32
_SHARED_WORD = 0.25

# Passages read out: the best, and others nearly as good from the same report.
_CLOSE = 0.08
_SPOKEN_LIMIT = 2

_PLAIN_WORDS = frozenset({
    "the", "and", "for", "about", "my", "your", "with", "that", "this", "from",
    "what", "did", "have", "any", "all", "our", "are", "was", "were", "has",
    "had", "been", "into", "its", "report", "reports", "research", "say", "said",
})

_SMALL_WORDS = frozenset({"my", "the", "a", "an", "your", "our", "any", "some", "all"})


def _chosen_reports(subject):
    """The reports a question is about: all of them, or those its subject names."""
    paths = report_files()

    if not subject or not paths:
        return paths

    from actions import semantic_memory

    labels = [describe_report(path)[0] for path in paths]
    named = semantic_memory.shares_words(subject, labels, _PLAIN_WORDS)
    chosen = [path for path, hit in zip(paths, named) if hit]

    if chosen:
        return chosen

    meaning = semantic_memory.similarities(subject, labels)

    if meaning:
        best = max(meaning)

        if best >= 0.5:
            return [path for path, score in zip(paths, meaning) if score >= best - 0.05]

    return []


def search(topic, subject=None):
    """Passages about [topic], best first, as (score, path, heading, passage)."""
    topic = " ".join(str(topic or "").split())

    if not topic:
        return []

    found = passages(_chosen_reports(subject))

    if not found:
        return []

    from actions import semantic_memory

    core = " ".join(word for word in topic.casefold().split() if word not in _SMALL_WORDS) or topic
    texts = [_encoded(heading, passage) for _path, heading, passage in found]

    try:
        meaning = semantic_memory.similarities(core, texts)
    except Exception as error:
        print(f"[JARVIS] searching reports by meaning failed, using words: {error}")
        meaning = None

    shared = semantic_memory.shares_words(core, texts, _PLAIN_WORDS)
    scored = []

    for index, (path, heading, passage) in enumerate(found):
        if meaning is None:
            score = 1.0 if shared[index] else 0.0
        else:
            score = meaning[index] + (_SHARED_WORD if shared[index] else 0.0)

        if score >= _MINIMUM:
            scored.append((score, path, heading, passage))

    scored.sort(key=lambda item: -item[0])

    return scored


# ---- what is said ------------------------------------------------------------

def _spoken_day(day, today):
    if day is None:
        return None

    if day == today:
        return "today"

    spoken = f"{day.day} {day.strftime('%B')}"

    return spoken if day.year == today.year else f"{spoken} {day.year}"


def _named(path, today):
    """"your ESP32 power report from 24 September", or "today's ... report"."""
    label, day = describe_report(path)

    if day == today:
        return f"today's {label} report"

    when = _spoken_day(day, today)

    return f"your {label} report" + (f" from {when}" if when else "")


def _yours(topic):
    swaps = {"my": "your", "mine": "yours", "me": "you", "i": "you", "myself": "yourself"}

    return " ".join(swaps.get(word.casefold(), word) for word in topic.split())


def answer(said, topic=None, mode=None, subject=None, today=None):
    """The spoken answer to a question about saved reports."""
    today = today or date.today()
    asked = question(said)

    if asked:
        mode = mode or asked[0]
        topic = topic or asked[1]
        subject = subject or asked[2]

    mode = mode or ("which" if re.match(r"^(?:which|what\s+reports|do\s+any)", _normalise(said)) else "say")
    topic = " ".join(str(topic or "").split())

    if not topic:
        return "What would you like me to look for in your reports, sir?"

    if not report_files():
        return "You have no saved reports yet, sir."

    if subject and not _chosen_reports(subject):
        return f"I can't find a report on {_yours(subject)}, sir."

    found = search(topic, subject)
    about = _yours(topic)

    if not found:
        where = f"your {_yours(subject)} report" if subject else "your reports"
        return f"I can't find anything about {about} in {where}, sir."

    if mode == "which":
        order = list(dict.fromkeys(path for _score, path, _heading, _passage in found))
        names = [describe_report(path) for path in order]
        spoken = [
            label + (f", {_spoken_day(day, today)}" if day else "")
            for label, day in names
        ]

        if len(spoken) == 1:
            return f"One report mentions {about}, sir: {spoken[0]}."

        listed = "; ".join(spoken[:-1]) + f"; and {spoken[-1]}"
        return f"{len(spoken)} reports mention {about}, sir: {listed}."

    best_score, best_path = found[0][0], found[0][1]
    chosen = [
        passage for score, path, _heading, passage in found
        if path == best_path and score >= best_score - _CLOSE
    ][:_SPOKEN_LIMIT]

    others = len({path for _score, path, _heading, _passage in found if path != best_path})
    also = ""

    if others == 1:
        also = " Another report mentions it too."
    elif others > 1:
        also = f" {others} other reports mention it too."

    source = _named(best_path, today)
    body = " ".join(passage.rstrip() if passage.rstrip().endswith((".", "!", "?")) else passage.rstrip() + "."
                    for passage in chosen)

    return f"From {source}, sir: {body}{also}"


# ---- indexing in the background --------------------------------------------

def index_quietly(delay=45.0, pause=0.2):
    """Encode report passages the vector store does not have yet."""
    time.sleep(delay)

    try:
        from actions import semantic_memory, vector_store

        texts = [_encoded(heading, passage) for _path, heading, passage in passages()]

        for start in range(0, len(texts), 128):
            if vector_store.vectors(texts[start:start + 128], semantic_memory._encode,
                                    semantic_memory.MODEL_ID) is None:
                return

            time.sleep(pause)

        if texts:
            print(f"[JARVIS] reports indexed for search ({len(texts)} passages)")
    except Exception as error:
        print(f"[JARVIS] could not index reports: {error}")
