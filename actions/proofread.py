"""Look for spelling mistakes in a file, entirely locally.

pyspellchecker ships a dictionary and is pure Python, so this costs nothing
and sends nothing anywhere. It reads plain text and Word documents through
the existing file layer, so the sandbox still applies.

Deliberately cautious: it reports *possible* mistakes. A dictionary knows
nothing about names, jargon, code or British spellings it happens to lack,
so the wording is always "possible", and fixing anything asks first.
"""

import os
import re

import phrases
from actions import files

try:
    from spellchecker import SpellChecker

    _AVAILABLE = True
except ImportError:
    print("[JARVIS] spell checking unavailable: pip install pyspellchecker")
    _AVAILABLE = False


# Words shorter than this are skipped: initials and abbreviations produce
# far more noise than genuine findings.
MIN_LENGTH = 4

# A very long file would take a while and produce an unusable list, so only
# this many words are examined.
MAX_WORDS = 60_000

# Reported at most, however many are found.
MAX_FINDINGS = 200

# Read aloud at most, since more than a few is unusable by voice.
SPOKEN_LIMIT = 4

_WORD = re.compile(r"[A-Za-z][A-Za-z'’]*")

# Anything matching these is not prose and should not be spell checked.
_SKIP = (
    re.compile(r"https?://\S+"),
    re.compile(r"\S+@\S+\.\S+"),
    re.compile(r"[A-Za-z]:\\\\\S+"),
    re.compile(r"\b[A-Z]{2,}\b"),
    re.compile(r"\b\w*\d\w*\b"),
    re.compile(r"`[^`]*`"),
)

_checker = None

# Words the user has told JARVIS to leave alone, kept in the JARVIS folder
# so they survive a restart and can be edited by hand.
IGNORE_FILE = "spelling-ignore.txt"


def available():
    return _AVAILABLE


def _spell():
    global _checker

    if _checker is None and _AVAILABLE:
        _checker = SpellChecker(language="en", distance=2)

    return _checker


def ignored_words():
    """Words to skip, from the ignore list in the JARVIS folder."""
    base = files.root()

    if not base:
        return set()

    path = os.path.join(base, IGNORE_FILE)

    if not os.path.exists(path):
        return set()

    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as handle:
            return {
                line.strip().casefold()
                for line in handle
                if line.strip() and not line.startswith("#")
            }

    except OSError:
        return set()


def ignore(word):
    """Add a word to the ignore list. Returns True on success."""
    word = (word or "").strip()

    if not word:
        return False

    base = files.root()

    if not base:
        return False

    path = os.path.join(base, IGNORE_FILE)

    try:
        fresh = not os.path.exists(path)

        with open(path, "a", encoding="utf-8") as handle:
            if fresh:
                handle.write(
                    "# Words JARVIS should not flag. One per line.\n"
                )

            handle.write(f"{word}\n")

        return True

    except OSError as error:
        print(f"[JARVIS] could not update the ignore list: {error}")
        return False


# Word separates table cells and rows with control characters rather than
# newlines, so a whole table arrives as one line. Left alone, every cell
# after the first looks like a capitalised word mid-sentence -- a name --
# and gets skipped, which is why "Seperate" was missed on screen but found
# in the file.
_CELL_MARKS = re.compile(r"[\x07\x0b\x0c\r\x1e\x1f]+")


def _split_cells(content):
    """Turn cell and row markers into real line breaks."""
    return _CELL_MARKS.sub("\n", content or "")


def _strip_noise(line):
    """Remove things that are not prose before checking."""
    for pattern in _SKIP:
        line = pattern.sub(" ", line)

    return line


def check(name):
    """Find possible spelling mistakes in a file.

    Returns (findings, total_words) where each finding is a dictionary with
    the line number, the word, and up to three suggestions. Returns
    (None, 0) when the file cannot be read.
    """
    content = files.read(name)

    if content is None:
        return None, 0

    return check_text(content)


def check_text(content):
    """Find possible spelling mistakes in a block of text.

    Used for files and for whatever is on screen, so both behave the same.
    """
    if not _AVAILABLE or content is None:
        return None, 0

    content = _split_cells(content)

    checker = _spell()

    if checker is None:
        return None, 0

    skip = ignored_words()

    findings = []
    seen = set()
    total = 0

    for number, line in enumerate(content.splitlines(), start=1):
        for match in _WORD.finditer(_strip_noise(line)):
            word = match.group(0)
            total += 1

            if total > MAX_WORDS:
                print(f"[JARVIS] only the first {MAX_WORDS} words were read")
                return findings, total

            if len(word) < MIN_LENGTH:
                continue

            lowered = word.casefold().replace("’", "'")

            if lowered in skip or lowered in seen:
                continue

            # A capitalised word mid-sentence is usually a name -- but not
            # when it opens a sentence, which is why "Thas" was being
            # missed. Look at what precedes it.
            if word[0].isupper():
                before = line[:match.start()].rstrip()

                if before and not before.endswith((".", "!", "?", ":", '"')):
                    continue

            if checker.known([lowered]):
                continue

            seen.add(lowered)

            # correction() picks the most likely by word frequency;
            # candidates come back as an unordered set, so sorting them
            # alphabetically was offering "seeling" for "speling".
            best = checker.correction(lowered)
            candidates = checker.candidates(lowered) or set()

            others = sorted(candidates - {lowered, best})
            suggestions = ([best] if best and best != lowered else []) + others
            suggestions = suggestions[:3]

            findings.append({
                "line": number,
                "word": word,
                "suggestions": suggestions,
            })

            if len(findings) >= MAX_FINDINGS:
                print(f"[JARVIS] stopping at {MAX_FINDINGS} findings")
                return findings, total

    return findings, total


def describe(name, findings, total):
    """A spoken summary of what was found."""
    label = files.spoken_name(files.safe_name(name) or name)

    if findings is None:
        return f"I couldn't read {label}, sir."

    if not findings:
        return (
            f"{label} looks clean, sir. "
            f"I checked {phrases.number(total)} words."
        )

    count = len(findings)
    word = "mistake" if count == 1 else "mistakes"

    return (
        f"I found {phrases.number(count)} possible spelling {word} in "
        f"{label}, sir, out of {phrases.number(total)} words."
    )


def spoken_list(findings, limit=SPOKEN_LIMIT):
    """The first few findings, read out.

    With many mistakes the count comes first, so you know how much you are
    not hearing before the reading starts.
    """
    if not findings:
        return "Nothing to report, sir."

    parts = []

    for finding in findings[:limit]:
        suggestion = finding["suggestions"][0] if finding["suggestions"] else None

        if suggestion:
            parts.append(
                f"line {finding['line']}, {finding['word']}, "
                f"perhaps {suggestion}"
            )
        else:
            parts.append(f"line {finding['line']}, {finding['word']}")

    listed = ". ".join(parts)

    if len(findings) <= limit:
        return f"{listed}, sir."

    return (
        f"There are {phrases.number(len(findings))}, sir. I'll read the "
        f"first {phrases.number(limit)}. {listed}. "
        "A report would give you the rest."
    )


def _matched_case(original, replacement):
    """Give the replacement the same capitalisation as the word it replaces.

    Without this, correcting "Thas" at the start of a sentence produced
    "that" and quietly introduced a different mistake.
    """
    if not original or not replacement:
        return replacement

    if original.isupper() and len(original) > 1:
        return replacement.upper()

    if original[0].isupper():
        return replacement[0].upper() + replacement[1:]

    return replacement


def corrected_text(content, findings):
    """The text with each flagged word replaced by its best suggestion.

    Used for text on screen, where JARVIS cannot edit the application
    directly: the corrected version goes on the clipboard for you to paste.
    """
    if not content:
        return ""

    for finding in findings or ():
        if not finding["suggestions"]:
            continue

        pattern = re.compile(rf"\b{re.escape(finding['word'])}\b")
        content = pattern.sub(
            _matched_case(finding["word"], finding["suggestions"][0]),
            content,
        )

    return content


def for_clipboard(content):
    """Tidy screen text so it pastes sensibly.

    Word marks the end of a cell and the end of a row with control
    characters. Pasted raw they are invisible rubbish, so cells become tabs
    and rows become line breaks -- which Word can turn back into a table
    with Insert, Table, Convert Text to Table.
    """
    if not content:
        return ""

    # Row ends first: two markers together mean end of row.
    tidied = re.sub(r"\x07\x07+", "\n", content)
    tidied = tidied.replace("\x07", "\t")
    tidied = re.sub(r"[\x0b\x0c\x1e\x1f]+", "\n", tidied)
    tidied = tidied.replace("\r\n", "\n").replace("\r", "\n")

    # Trailing tabs at the end of a row serve no purpose once pasted.
    tidied = re.sub(r"\t+\n", "\n", tidied)

    return tidied.strip()


def report(name, findings, total):
    """Write the findings to a file. Returns the path, or None."""
    label = files.safe_name(name) or name
    stem = os.path.splitext(label)[0]

    lines = [
        f"Spelling report for {label}",
        f"{total} words checked, {len(findings or ())} possible mistakes",
        "",
    ]

    for finding in findings or ():
        suggestions = ", ".join(finding["suggestions"]) or "no suggestion"
        lines.append(
            f"line {finding['line']:>5}  {finding['word']:<24} {suggestions}"
        )

    lines.extend([
        "",
        "Words here that are correct can be added to "
        f"{IGNORE_FILE} so they are not flagged again.",
    ])

    title = stem if stem.casefold().endswith(
        "spelling") else f"{stem} spelling"

    return files.write(
        f"{title} report", "\n".join(lines), default_suffix=".txt"
    )


def _fix_docx(path, findings):
    """Correct words inside a Word document, keeping its formatting.

    Replacements happen run by run rather than paragraph by paragraph:
    setting a paragraph's text wholesale would discard the bold, italics
    and styling of everything in it.
    """
    try:
        from docx import Document
    except ImportError:
        print("[JARVIS] python-docx is not installed")
        return None

    try:
        document = Document(path)
    except Exception as error:
        print(f"[JARVIS] could not open {path}: {error}")
        return None

    changed = 0

    def mend(runs):
        nonlocal changed

        for run in runs:
            if not run.text:
                continue

            updated = run.text

            for finding in findings:
                pattern = re.compile(rf"\b{re.escape(finding['word'])}\b")
                updated, count = pattern.subn(
                    _matched_case(finding["word"], finding["suggestions"][0]),
                    updated,
                )
                changed += count

            if updated != run.text:
                run.text = updated

    for paragraph in document.paragraphs:
        mend(paragraph.runs)

    # Text inside tables is easy to forget and common in real documents.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    mend(paragraph.runs)

    if not changed:
        # A word split across two runs is not visible to this approach.
        print("[JARVIS] nothing matched inside the document")
        return 0

    try:
        document.save(path)
    except Exception as error:
        print(f"[JARVIS] could not save {path}: {error}")
        return None

    print(f"[JARVIS] corrected {changed} word(s) in {path}")

    return changed


def apply_fixes(name, findings):
    """Replace each flagged word with its first suggestion.

    Only words with a suggestion are touched, and each is replaced whole,
    so a short word cannot corrupt a longer one containing it. Returns the
    number of replacements, or None if the file could not be written.
    """
    usable = [f for f in findings or () if f["suggestions"]]

    if not usable:
        return 0

    content = files.read(name)

    if content is None:
        return None

    changed = 0

    for finding in usable:
        pattern = re.compile(
            rf"\b{re.escape(finding['word'])}\b"
        )
        content, count = pattern.subn(
            _matched_case(finding["word"], finding["suggestions"][0]),
            content,
        )
        changed += count

    path = files.find_existing(name)

    if not path:
        return None

    suffix = os.path.splitext(path)[1].casefold()

    if suffix == ".docx":
        return _fix_docx(path, usable)

    try:
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(content)

    except OSError as error:
        print(f"[JARVIS] could not write {path}: {error}")
        return None

    print(f"[JARVIS] corrected {changed} word(s) in {path}")

    return changed
