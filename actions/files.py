"""File operations, confined to a single folder.

Everything JARVIS reads or writes lives under ~/JARVIS. A misheard command
should never be able to touch anything else, so every path is resolved and
checked against that root before use.
"""

import os
import re
import shutil

# Optional formats. Missing libraries are reported rather than crashing.
try:
    from docx import Document

    _DOCX = True
except ImportError:
    _DOCX = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    _PDF = True
except ImportError:
    _PDF = False


FOLDER_NAME = "JARVIS"

TEXT_SUFFIXES = (".txt", ".md", ".csv", ".log", ".json")

# Spoken format words mapped to a file extension.
FORMATS = {
    "text": ".txt", "text file": ".txt", "txt": ".txt", "plain text": ".txt",
    "markdown": ".md", "md": ".md",
    "word": ".docx", "word document": ".docx", "doc": ".docx",
    "docx": ".docx", "word doc": ".docx",
    "pdf": ".pdf", "pdf document": ".pdf",
    "csv": ".csv", "spreadsheet": ".csv",
}

_ILLEGAL = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

# Reading a whole file aloud is unusable, so only this much is spoken.
_SPEAK_LIMIT = 400


def root():
    """The one folder JARVIS may touch. Created if missing."""
    path = os.path.join(os.path.expanduser("~"), FOLDER_NAME)

    try:
        os.makedirs(path, exist_ok=True)
    except OSError as error:
        print(f"[JARVIS] could not create {path}: {error}")
        return None

    return path


def safe_name(name, default_suffix=".txt"):
    """Turn spoken words into a filename, or None if nothing usable remains."""
    cleaned = _ILLEGAL.sub("", (name or "").strip())
    cleaned = " ".join(cleaned.split())
    cleaned = cleaned.strip(". ")

    if not cleaned:
        return None

    stem, suffix = os.path.splitext(cleaned)

    if not suffix:
        suffix = default_suffix

    return f"{stem}{suffix}"


def resolve(name, default_suffix=".txt"):
    """Full path inside the folder, or None if it would escape it."""
    base = root()
    filename = safe_name(name, default_suffix)

    if not base or not filename:
        return None

    path = os.path.abspath(os.path.join(base, filename))

    # The decisive check: anything that resolves outside the root is refused.
    if os.path.commonpath([path, os.path.abspath(base)]) != os.path.abspath(base):
        print(f"[JARVIS] refusing a path outside {base}")
        return None

    return path


def exists(name, default_suffix=".txt"):
    """True when a file of this name exists, whatever its extension."""
    path = resolve(name, default_suffix)

    if path and os.path.exists(path):
        return True

    return find_existing(name) is not None


def unique_path(path):
    """A path that does not exist yet, adding (1), (2) and so on."""
    if not os.path.exists(path):
        return path

    stem, suffix = os.path.splitext(path)

    for index in range(1, 500):
        candidate = f"{stem} ({index}){suffix}"

        if not os.path.exists(candidate):
            return candidate

    return None


def write(name, content="", default_suffix=".txt", overwrite=False):
    """Create a file with optional content. Returns the path, or None."""
    path = resolve(name, default_suffix)

    if not path:
        return None

    if os.path.exists(path) and not overwrite:
        path = unique_path(path)

        if not path:
            return None

    suffix = os.path.splitext(path)[1].casefold()

    try:
        if suffix == ".docx":
            if not _DOCX:
                print("[JARVIS] python-docx is not installed")
                return None

            document = Document()

            for line in (content or "").split("\n"):
                document.add_paragraph(line)

            document.save(path)

        elif suffix == ".pdf":
            if not _PDF:
                print("[JARVIS] reportlab is not installed")
                return None

            doc = SimpleDocTemplate(path, pagesize=A4)
            styles = getSampleStyleSheet()
            flow = []

            for line in (content or "").split("\n"):
                if line.strip():
                    flow.append(Paragraph(line, styles["Normal"]))
                else:
                    flow.append(Spacer(1, 10))

            doc.build(flow or [Spacer(1, 10)])

        else:
            with open(path, "w", encoding="utf-8") as handle:
                handle.write(content or "")

    except Exception as error:
        print(f"[JARVIS] could not write {path}: {error}")
        return None

    print(f"[JARVIS] wrote {path}")

    return path


KNOWN_SUFFIXES = (".txt", ".md", ".docx", ".pdf", ".csv", ".log", ".json")


def find_existing(name):
    """Path of a file matching this spoken name, whatever its extension.

    "business" should find business.docx rather than inventing business.txt.
    """
    exact = resolve(name)

    if exact and os.path.exists(exact):
        return exact

    stem = os.path.splitext(safe_name(name) or "")[0]

    if not stem:
        return None

    for suffix in KNOWN_SUFFIXES:
        candidate = resolve(f"{stem}{suffix}")

        if candidate and os.path.exists(candidate):
            return candidate

    return None


def append(name, content, default_suffix=".txt"):
    """Add a line to an existing file, creating a text file if none exists."""
    path = find_existing(name) or resolve(name, default_suffix)

    if not path:
        return None

    suffix = os.path.splitext(path)[1].casefold()

    if suffix == ".docx":
        return _append_docx(path, content)

    if suffix not in TEXT_SUFFIXES:
        print(f"[JARVIS] cannot append to {suffix} files")
        return None

    try:
        # Existing content may not end with a newline, which would run the
        # two lines together.
        needs_break = os.path.exists(path) and os.path.getsize(path) > 0

        if needs_break:
            with open(path, "rb") as handle:
                handle.seek(-1, os.SEEK_END)
                needs_break = handle.read(1) not in (b"\n", b"\r")

        with open(path, "a", encoding="utf-8") as handle:
            if needs_break:
                handle.write("\n")

            handle.write(f"{content}\n")

    except OSError as error:
        print(f"[JARVIS] could not append to {path}: {error}")
        return None

    return path


def _append_docx(path, content):
    """Add a paragraph to an existing Word document."""
    if not _DOCX:
        print("[JARVIS] python-docx is not installed")
        return None

    try:
        document = Document(path)
        document.add_paragraph(content)
        document.save(path)

    except Exception as error:
        print(f"[JARVIS] could not append to {path}: {error}")
        return None

    return path


def _remove_docx_lines(path, wanted):
    """Delete matching paragraphs from a Word document."""
    if not _DOCX:
        print("[JARVIS] python-docx is not installed")
        return None

    try:
        document = Document(path)

        removed = 0
        remaining = 0

        for paragraph in list(document.paragraphs):
            if wanted in paragraph.text.casefold():
                # A paragraph is removed by detaching its XML element; there
                # is no delete method on the object itself.
                element = paragraph._element
                element.getparent().remove(element)
                removed += 1
            elif paragraph.text.strip():
                remaining += 1

        if removed:
            document.save(path)
            print(f"[JARVIS] removed {removed} paragraph(s) from {path}")

        return removed, remaining

    except Exception as error:
        print(f"[JARVIS] could not update {path}: {error}")
        return None


def remove_line(name, text):
    """Delete lines containing this text. Returns (removed, remaining).

    Returns None when the file exists but cannot be edited, so the caller can
    say why rather than claiming the text was not found.
    """
    wanted = " ".join((text or "").split()).casefold()

    if not wanted:
        return 0, 0

    path = find_existing(name)

    if not path:
        return None

    suffix = os.path.splitext(path)[1].casefold()

    if suffix == ".docx":
        return _remove_docx_lines(path, wanted)

    if suffix not in TEXT_SUFFIXES:
        print(f"[JARVIS] cannot edit {suffix} files")
        return None

    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as handle:
            lines = handle.read().splitlines()

    except OSError as error:
        print(f"[JARVIS] could not read {path}: {error}")
        return 0, 0

    kept = [line for line in lines if wanted not in line.casefold()]
    removed = len(lines) - len(kept)

    if not removed:
        return 0, len([line for line in lines if line.strip()])

    try:
        with open(path, "w", encoding="utf-8") as handle:
            for line in kept:
                handle.write(f"{line}\n")

    except OSError as error:
        print(f"[JARVIS] could not update {path}: {error}")
        return 0, len(lines)

    print(f"[JARVIS] removed {removed} line(s) from {path}")

    return removed, len([line for line in kept if line.strip()])


def describe_removal(name, text):
    """Remove lines from a file and report what happened."""
    path = find_existing(name)
    label = spoken_name(safe_name(name) or name)

    if not path:
        return f"I couldn't find a file called {name}, sir."

    result = remove_line(name, text)

    if result is None:
        suffix = os.path.splitext(path)[1].lstrip(
            ".").upper() or "that kind of"

        return f"I can't edit {suffix} files, sir."

    removed, remaining = result

    if not removed:
        return f"I couldn't find {text} in {label}, sir."

    if removed == 1:
        return f"Removed {text} from {label}, sir. {remaining} lines left."

    return f"Removed {removed} lines mentioning {text} from {label}, sir."


def read(name):
    """Return a file's text, or None if it cannot be read."""
    path = find_existing(name)

    if not path:
        return None

    suffix = os.path.splitext(path)[1].casefold()

    try:
        if suffix == ".docx":
            if not _DOCX:
                print("[JARVIS] python-docx is not installed")
                return None

            document = Document(path)

            return "\n".join(p.text for p in document.paragraphs)

        if suffix in TEXT_SUFFIXES:
            with open(path, "r", encoding="utf-8", errors="ignore") as handle:
                return handle.read()

        print(f"[JARVIS] cannot read {suffix} files")
        return None

    except Exception as error:
        print(f"[JARVIS] could not read {path}: {error}")
        return None


def copy(source, destination=None):
    """Copy a file within the folder. Returns the new path, or None."""
    source_path = find_existing(source)

    if not source_path:
        print(f"[JARVIS] no file named {source!r}")
        return None

    if destination:
        target = resolve(destination, os.path.splitext(source_path)[1])
    else:
        stem, suffix = os.path.splitext(os.path.basename(source_path))
        target = resolve(f"{stem} copy{suffix}")

    if not target:
        return None

    target = unique_path(target)

    if not target:
        return None

    try:
        shutil.copy2(source_path, target)
    except OSError as error:
        print(f"[JARVIS] could not copy: {error}")
        return None

    print(f"[JARVIS] copied to {target}")

    return target


def listing():
    """Filenames in the folder, newest first."""
    base = root()

    if not base:
        return []

    try:
        entries = [
            entry for entry in os.listdir(base)
            if os.path.isfile(os.path.join(base, entry))
        ]
    except OSError:
        return []

    entries.sort(
        key=lambda name: os.path.getmtime(os.path.join(base, name)),
        reverse=True,
    )

    return entries


_SCREENSHOT = re.compile(r"^jarvis-\d{4}-\d{2}-\d{2}-\d{6}$", re.IGNORECASE)

_TYPE_WORDS = {
    ".pdf": "PDF",
    ".docx": "Word document",
    ".doc": "Word document",
    ".csv": "spreadsheet",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".md": "markdown file",
    ".txt": "",
    ".log": "log",
    ".json": "JSON file",
}


def spoken_name(filename):
    """A filename a voice can say quickly.

    Timestamped names are read digit by digit otherwise, which takes about
    fifteen seconds each.
    """
    stem, suffix = os.path.splitext(filename)
    suffix = suffix.casefold()

    if _SCREENSHOT.match(stem):
        return "a screenshot"

    kind = _TYPE_WORDS.get(suffix, suffix.lstrip(".") or "")

    # Underscores and hyphens are spoken aloud, so replace them.
    spoken = stem.replace("_", " ").replace("-", " ")
    spoken = " ".join(spoken.split())

    return f"{spoken} {kind}".strip() if kind else spoken


def describe_listing(limit=4, names=True):
    """Spoken summary of what is in the folder.

    The true total is always given. Names are read for the most recent few,
    since reading fifty filenames aloud would be unbearable.
    """
    entries = listing()

    if not entries:
        return "Your JARVIS folder is empty, sir."

    if not names:
        if len(entries) == 1:
            return "One file, sir."

        return f"You have {len(entries)} files, sir."

    if len(entries) == 1:
        return f"One file, sir: {spoken_name(entries[0])}."

    shown = [spoken_name(entry) for entry in entries[:limit]]
    listed = ", ".join(shown[:-1]) + f", and {shown[-1]}"

    if len(entries) <= limit:
        return f"You have {len(entries)} files, sir: {listed}."

    return (
        f"You have {len(entries)} files, sir. I won't list them all, but "
        f"your {len(shown)} most recent are {listed}."
    )


def describe_listing_count():
    """Just the number, for "how many files do I have"."""
    return describe_listing(names=False)


def describe_listing_named(limit=4):
    """The full version, with filenames read out."""
    return describe_listing(limit=limit, names=True)


# Signals that on their own strongly imply code.
_CODE_STRONG = (
    "function ", "const ", "=>", "def ", "class ", "#include", "<?php",
    "console.log", "return ", "import ", "async ", "await ", "public ",
    "private ", "};", "();", "()", "SELECT ", "INSERT ", "UPDATE ",
)

# Weaker signals, which need company to count.
_CODE_WEAK = (
    "let ", "var ", "if (", "for (", "while (", "print(", "{", "}",
    "==", "!=", "&&", "||", "self.", "this.", "FROM ", "WHERE ",
)

# Above this share of symbols the file is treated as code regardless.
_SYMBOL_SHARE = 0.12


def looks_like_code(text):
    """True when reading this aloud would be gibberish."""
    sample = text[:2000]
    stripped = sample.strip()

    if not stripped:
        return False

    # JSON has few keywords but unmistakable structure.
    if stripped[0] in "{[" and stripped[-1] in "}]" and '":' in sample:
        return True

    strong = sum(1 for hint in _CODE_STRONG if hint in sample)
    weak = sum(1 for hint in _CODE_WEAK if hint in sample)

    # One strong signal plus any support, or several weak ones together.
    if strong >= 2 or (strong and weak) or weak >= 4:
        return True

    symbols = sum(1 for ch in sample if ch in "{}[]()<>;=/\\|*&^%$#@~`_")

    return symbols / len(sample) > _SYMBOL_SHARE


def describe_read(name):
    """Spoken rendering of a file's contents."""
    content = read(name)
    label = spoken_name(safe_name(name) or name)

    if content is None:
        return f"I couldn't find a file called {label}, sir."

    tidied = " ".join(content.split())

    if not tidied:
        return f"{label} is empty, sir."

    if looks_like_code(content):
        lines = len(content.splitlines())
        words = len(tidied.split())

        return (
            f"{label} looks like code, sir: {lines} lines, "
            f"{words} words. I won't read it aloud."
        )

    if len(tidied) <= _SPEAK_LIMIT:
        return f"{label} says: {tidied}"

    words = len(tidied.split())
    opening = tidied[:_SPEAK_LIMIT].rsplit(" ", 1)[0]

    return f"{label} holds {words} words, sir. It begins: {opening}..."
