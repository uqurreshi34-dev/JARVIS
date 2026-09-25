"""How well JARVIS finds things: retrieval measured, not guessed.

Every search JARVIS does by meaning -- memory, notes, the command log,
saved reports -- is scored here, so a change to any of them is proven
better or worse rather than assumed. Everything is local: no model call,
no network, no cost.

    python tools/rag_eval.py            the benchmark, compared with the saved baseline
    python tools/rag_eval.py --update   the same, then saves these scores as the baseline
    python tools/rag_eval.py --mine     your own questions and reports, read-only

The benchmark runs in a sandboxed JARVIS folder on fixed data: notes,
memories, a 600-command log and reports written the way JARVIS writes
them. For each store it measures:

    first right   the right thing comes first (hit@1)
    in top three  the right thing is among the first three (hit@3)
    says nothing  a question nothing is about finds nothing
    precision / recall   for the log, whether a count includes exactly
                  the commands about the topic

and how long a search takes. A score that falls more than 0.02 below the
baseline in tools/rag_baseline.json fails the run.

--mine reads rag-questions.txt in your JARVIS folder, one question a line:

    notes   | what did I note about the heating       | boiler
    log     | when did I last ask about the weather   | weather
    reports | what did my report say about villa park | Villa Park
    memory  | when do I train                         | monday
    notes   | what did I note about football          | -

The last column is text the answer must contain, or "-" when nothing
should be found. It also reports how well your saved reports cite their
sources. None of your files is changed; the only thing written is the
vector cache (jarvis-vectors.sqlite), exactly as when JARVIS searches.
"""

import argparse
import json
import os
import random
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.environ.setdefault("GROQ_API_KEY", "rag-eval")

BASELINE = ROOT / "tools" / "rag_baseline.json"
QUESTIONS_NAME = "rag-questions.txt"

# How far a score may fall below the baseline before the run fails.
TOLERANCE = 0.02


# ---- the benchmark data ------------------------------------------------------

MEMORIES = [
    "name: avid coder",
    "gym days: monday, wednesday and friday",
    "default project: jarvis",
    "reply length: short",
    "I drive a blue Honda Civic",
    "my favourite football team is Aston Villa",
    "my dentist is Dr Patel on Soho Road",
    "I usually order a flat white",
    "my cousin writes an ONNX to C compiler called FEMTO-ML",
    "I am learning to paint with watercolours",
    "my phone is a Samsung Galaxy A17",
    "I work from my room in a shared house",
]

# question -> text the first memory found must contain
MEMORY_QUESTIONS = {
    "what are my gym days": "gym days",
    "when do I train": "gym days",
    "which days do I work out": "gym days",
    "what project am I working on": "default project",
    "what car do I drive": "honda civic",
    "which team do I support": "aston villa",
    "who is my dentist": "dr patel",
    "what coffee do I like": "flat white",
    "what is my cousin building": "femto-ml",
    "what hobby am I learning": "watercolours",
    "what phone do I have": "samsung",
    "where do I work": "shared house",
}

MEMORY_NOTHING = ["what is my blood type", "what is my favourite film", "where did I go on holiday"]

NOTES = [
    "call the plumber about the boiler on tuesday",
    "buy milk, eggs and bread",
    "the wifi password for the guest network is on the back of the router",
    "dentist appointment moved to the 14th",
    "idea: add a radar sweep animation to the HUD",
    "renew car insurance before october",
    "ask cousin about the femto ml core question",
    "pay council tax",
    "book train tickets to manchester for the wedding",
    "sensor board needs a 10k resistor for the DHT22",
    "remember to back up the askfiles keystore",
    "gym on monday wednesday friday",
    "birthday present for mum",
    "check render logs for the throttle warnings",
    "look into tailscale magic dns on android",
]

NOTE_QUESTIONS = {
    "the boiler": 0, "heating": 0, "plumber": 0, "groceries": 1, "shopping": 1, "wifi": 2,
    "internet password": 2, "passwords": 2, "teeth": 3, "appointments": 3, "hud ideas": 4,
    "insurance": 5, "my car": 5, "my cousin": 6, "tax": 7, "trains": 8, "the wedding": 8,
    "resistor": 9, "dht22": 9, "keystore": 10, "backups": 10, "ask files": 10, "exercise": 11,
    "workout": 11, "my mum": 12, "gifts": 12, "render": 13, "tailscale": 14, "vpn": 14,
}

NOTE_NOTHING = ["quantum physics", "football", "the cat", "crypto prices", "the weather", "my phone", "my plans"]

LOG_GROUPS = {
    "bitcoin": ["whats the bitcoin price", "how is bitcoin doing", "bitcoin price please",
                "give me a market report on bitcoin"],
    "chrome": ["open chrome", "close chrome", "open google chrome"],
    "weather": ["whats the weather like", "will it rain tomorrow", "whats the weather in london"],
    "github": ["what are my git hub issues", "list my github repositories", "any new pull requests on github"],
    "music": ["play the next song", "pause the music", "close spotify"],
    "other": ["what time is it", "take a screenshot", "whats in my clipboard", "whats on my calendar",
              "how much battery do i have", "show me the news", "what planes are overhead",
              "set a timer for ten minutes", "recite surah al fatiha", "turn the volume up",
              "model a chair in blender", "proofread this document", "read my notes"],
}

# topic -> the group every counted command must belong to
LOG_QUESTIONS = {
    "bitcoin": "bitcoin", "crypto": "bitcoin", "chrome": "chrome", "the browser": "chrome",
    "the weather": "weather", "github": "github", "music": "music",
}

LOG_NOTHING = ["football", "my code", "quantum physics", "the dentist"]

REPORTS = {
    ("Reports", "ESP32 power consumption research report 2026-09-20-101500.docx"): """# ESP32 Power Consumption
## Executive summary
The ESP32-WROOM-32 draws between 160 and 260 milliamps while transmitting over Wi-Fi [1].
## Active and transmit current
During Wi-Fi transmission peaks can briefly exceed 300 milliamps, so the supply must deliver short bursts without the voltage sagging below 3.0 volts. Brownout resets are the most common symptom of an undersized regulator [1].
## Deep sleep and wake sources
Deep sleep keeps only the RTC controller and RTC memory powered. The board can wake on a timer, a touch pad, or an external pin such as a PIR motion sensor output [2].
## Choosing a battery
A 2000 mAh lithium cell running a sensor that wakes every five minutes lasts roughly three to five weeks on a bare module [2].
## Sources
1. Espressif datasheet - https://www.espressif.com""",
    ("", "Bitcoin vs Ethereum research report 2026-09-22-201000.docx"): """# Bitcoin vs Ethereum
## Summary
Bitcoin is designed as a store of value with a fixed supply of 21 million coins, while Ethereum is a programmable platform [1].
## Supply and issuance
Bitcoin's issuance halves roughly every four years [1].
## Energy use
Bitcoin's proof of work mining consumes a large amount of electricity. Ethereum's switch to proof of stake in 2022 cut its energy use by more than 99 percent [2].
## Fees and speed
Ethereum transaction fees rise sharply when the network is congested [2].
## Sources
1. Bitcoin whitepaper - https://bitcoin.org""",
    ("Reports", "Dash cams research report 2026-08-30-090000.docx"): """# Dash Cam Comparison
**Summary**
For most drivers a front and rear dual-channel dash cam recording at 1440p offers the best balance [1].
**Storage and memory cards**
Dash cams overwrite old footage in a loop, so a high-endurance microSD card rated for continuous writing is essential [1].
**Insurance**
Some UK insurers offer a small discount for fitted dash cams [1].
**Sources**
1. Which? - https://www.which.co.uk""",
    ("Football", "Aston Villa Football Club vs Birmingham City Football Club research report 2026-09-16-120000.docx"): """# Aston Villa FC vs. Birmingham City FC - Comparative Briefing
*Prepared for filing under: Football*
---
## Executive Summary
This report was commissioned as a head-to-head comparison of Aston Villa Football Club and Birmingham City Football Club.
Aston Villa are the more successful club, with seven league titles and the 1982 European Cup [1].
## Aston Villa
Villa Park holds around 42,000 spectators and the club returned to European football in 2024 [2].
## Birmingham City
Birmingham City play at St Andrew's and were promoted back to the Championship in 2025 [3].
## Sources
1. Wikipedia - https://en.wikipedia.org""",
}

# question topic -> (report name starts with, section heading)
REPORT_QUESTIONS = {
    "deep sleep": ("ESP32", "Deep sleep and wake sources"),
    "battery life": ("ESP32", "Choosing a battery"),
    "how long will a battery last": ("ESP32", "Choosing a battery"),
    "brownouts": ("ESP32", "Active and transmit current"),
    "wifi current": ("ESP32", "Active and transmit current"),
    "energy use": ("Bitcoin", "Energy use"),
    "electricity": ("Bitcoin", "Energy use"),
    "halving": ("Bitcoin", "Supply and issuance"),
    "fees": ("Bitcoin", "Fees and speed"),
    "memory cards": ("Dash", "Storage and memory cards"),
    "sd cards": ("Dash", "Storage and memory cards"),
    "insurance": ("Dash", "Insurance"),
    "villa park": ("Aston Villa", "Aston Villa"),
    "st andrews": ("Aston Villa", "Birmingham City"),
}

REPORT_NOTHING = ["recipes", "the weather", "my holiday", "quantum computing", "javascript", "tax returns"]


# ---- measuring ---------------------------------------------------------------

class Suite:
    """Scores for one store, and the questions it got wrong."""

    def __init__(self, name):
        self.name = name
        self.scores = {}
        self.wrong = []
        self.seconds = []

    def timed(self, work):
        started = time.perf_counter()
        result = work()
        self.seconds.append(time.perf_counter() - started)
        return result

    def rate(self, metric, hits, total):
        self.scores[metric] = round(hits / total, 3) if total else None

    def milliseconds(self):
        # The first search encodes everything once; the rest are what you feel.
        later = self.seconds[1:] or self.seconds
        return round(1000 * sum(later) / len(later)) if later else None


def _memory_suite():
    from actions import memory, semantic_memory

    suite = Suite("memory")

    with open(memory._path(), "w", encoding="utf-8") as handle:
        handle.write("\n".join(MEMORIES) + "\n")

    semantic_memory.install()
    semantic_memory.clear_cache()

    def found(question):
        summary = memory.relevant_summary(question, limit=3)
        return [line[2:].casefold() for line in summary.splitlines() if line.startswith("- ")]

    first = top = 0

    for question, expected in MEMORY_QUESTIONS.items():
        lines = suite.timed(lambda: found(question))
        first += bool(lines) and expected in lines[0]
        top += any(expected in line for line in lines[:3])

        if not (lines and expected in lines[0]):
            suite.wrong.append(f"{question!r} -> {lines[:1] or 'nothing'} (wanted {expected!r})")

    quiet = 0

    for question in MEMORY_NOTHING:
        lines = suite.timed(lambda: found(question))
        quiet += not lines

        if lines:
            suite.wrong.append(f"{question!r} found {lines[0]!r} (wanted nothing)")

    suite.rate("first right", first, len(MEMORY_QUESTIONS))
    suite.rate("in top three", top, len(MEMORY_QUESTIONS))
    suite.rate("says nothing", quiet, len(MEMORY_NOTHING))

    return suite


def _notes_suite():
    from actions import notes

    suite = Suite("notes")
    today = date.today()

    with open(notes._notes_path(), "w", encoding="utf-8") as handle:
        for offset, text in enumerate(NOTES):
            handle.write(f"[{(today - timedelta(days=offset)).isoformat()} 09:30] {text}\n")

    first = top = 0

    for topic, index in NOTE_QUESTIONS.items():
        results = [text for _day, text in suite.timed(lambda: notes.search(topic, limit=3))]
        first += bool(results) and results[0] == NOTES[index]
        top += NOTES[index] in results[:3]

        if not (results and results[0] == NOTES[index]):
            suite.wrong.append(f"{topic!r} -> {results[:1] or 'nothing'} (wanted {NOTES[index]!r})")

    quiet = 0

    for topic in NOTE_NOTHING:
        results = suite.timed(lambda: notes.search(topic))
        quiet += not results

        if results:
            suite.wrong.append(f"{topic!r} found {results[0][1]!r} (wanted nothing)")

    suite.rate("first right", first, len(NOTE_QUESTIONS))
    suite.rate("in top three", top, len(NOTE_QUESTIONS))
    suite.rate("says nothing", quiet, len(NOTE_NOTHING))

    return suite


def _log_suite():
    from actions import files, log_search

    suite = Suite("log")
    now = datetime(2026, 9, 25, 20, 0)
    randomness = random.Random(7)
    log = []

    for step in range(600):
        group = randomness.choice(list(LOG_GROUPS))
        log.append((now - timedelta(minutes=37 * step), randomness.choice(LOG_GROUPS[group])))

    log.sort()

    with open(os.path.join(files.root(), "jarvis-log.txt"), "w", encoding="utf-8") as handle:
        for moment, said in log:
            handle.write(f"{moment:%Y-%m-%d %H:%M:%S}  command          {said!r} -> x (local)\n")

    precision = []
    recall = []

    for topic, group in LOG_QUESTIONS.items():
        found = suite.timed(lambda: log_search.search(topic, today=now.date()))
        wanted = [said for _moment, said in log if said in LOG_GROUPS[group]]
        right = [said for _moment, said, _intent in found if said in LOG_GROUPS[group]]
        precision.append(len(right) / len(found) if found else 0.0)
        recall.append(len(right) / len(wanted) if wanted else 1.0)

        if len(right) != len(found) or len(right) != len(wanted):
            extra = sorted({said for _m, said, _i in found} - set(LOG_GROUPS[group]))
            missed = sorted(set(LOG_GROUPS[group]) - {said for _m, said, _i in found})
            suite.wrong.append(f"{topic!r}: counted {len(found)} of {len(wanted)}; wrong {extra}, missed {missed}")

    quiet = 0

    for topic in LOG_NOTHING:
        found = suite.timed(lambda: log_search.search(topic, today=now.date()))
        quiet += not found

        if found:
            suite.wrong.append(f"{topic!r} counted {len(found)} (wanted none)")

    suite.scores["precision"] = round(sum(precision) / len(precision), 3)
    suite.scores["recall"] = round(sum(recall) / len(recall), 3)
    suite.rate("says nothing", quiet, len(LOG_NOTHING))

    return suite


def _reports_suite():
    from actions import files, report_search

    suite = Suite("reports")

    try:
        from docx import Document
    except ImportError:
        suite.wrong.append("python-docx is not installed; reports were not measured")
        return suite

    for (subfolder, name), text in REPORTS.items():
        where = os.path.join(files.root(), subfolder) if subfolder else files.root()
        os.makedirs(where, exist_ok=True)
        document = Document()

        for line in text.split("\n"):
            document.add_paragraph(line)

        document.save(os.path.join(where, name))

    first = top = 0

    for topic, (report, heading) in REPORT_QUESTIONS.items():
        results = suite.timed(lambda: report_search.search(topic))
        hits = [os.path.basename(path).startswith(report) and found_heading == heading
                for _score, path, found_heading, _passage in results[:3]]
        first += bool(hits) and hits[0]
        top += any(hits)

        if not (hits and hits[0]):
            got = (os.path.basename(results[0][1])[:20], results[0][2]) if results else "nothing"
            suite.wrong.append(f"{topic!r} -> {got} (wanted {report}, {heading!r})")

    quiet = 0

    for topic in REPORT_NOTHING:
        results = suite.timed(lambda: report_search.search(topic))
        quiet += not results

        if results:
            suite.wrong.append(f"{topic!r} found {results[0][2]!r} (wanted nothing)")

    suite.rate("first right", first, len(REPORT_QUESTIONS))
    suite.rate("in top three", top, len(REPORT_QUESTIONS))
    suite.rate("says nothing", quiet, len(REPORT_NOTHING))

    return suite


def benchmark():
    from tools import sandbox

    sandbox.activate()

    from actions import semantic_memory

    if semantic_memory._encode(["probe"]) is None:
        print("The local model is unavailable, so meaning-based search cannot be measured.")
        sys.exit(1)

    return [_memory_suite(), _notes_suite(), _log_suite(), _reports_suite()]


# ---- your own questions ------------------------------------------------------------

def mine():
    """Your questions from rag-questions.txt, and your reports' citations. Read-only."""
    from actions import files, log_search, memory, notes, report_search, semantic_memory

    semantic_memory.install()
    path = os.path.join(files.root() or "", QUESTIONS_NAME)

    if not os.path.exists(path):
        print(f"No {QUESTIONS_NAME} in your JARVIS folder yet. Create it with one question a line:\n")
        print("    notes   | what did I note about the heating       | boiler")
        print("    reports | what did my report say about villa park | Villa Park")
        print("    log     | when did I last ask about the weather   | weather")
        print("    memory  | when do I train                         | monday")
        print("    notes   | what did I note about football          | -\n")
        print('The last column is text the answer must contain, or "-" when nothing should be found.')
    else:
        _run_questions(path, files, log_search, memory, notes, report_search)

    _report_citations(report_search)


def _run_questions(path, files, log_search, memory, notes, report_search):
    passed = total = 0

    with open(path, encoding="utf-8") as handle:
        lines = [line.strip() for line in handle if line.strip() and not line.lstrip().startswith("#")]

    for number, line in enumerate(lines, start=1):
        parts = [part.strip() for part in line.split("|")]

        if len(parts) != 3 or parts[0].casefold() not in ("notes", "log", "reports", "memory"):
            print(f"  line {number}: not understood, skipped: {line!r}")
            continue

        store, question, wanted = parts[0].casefold(), parts[1], parts[2]
        nothing = wanted == "-"
        total += 1

        if store == "notes":
            topic = notes.search_topic(" ".join(re.sub(r"[^\w\s]", " ", question.casefold()).split())) or question
            found = bool(notes.search(topic))
            answer = notes.describe_search(topic)
        elif store == "log":
            asked = log_search.question(question)
            topic = asked[1] if asked else question
            found = bool(log_search.search(topic))
            answer = log_search.answer(question, topic=topic)
        elif store == "reports":
            asked = report_search.question(question)
            topic = asked[1] if asked else question
            found = bool(report_search.search(topic, asked[2] if asked else None))
            answer = report_search.answer(question, topic=topic)
        else:
            summary = memory.relevant_summary(question, limit=3)
            found = any(line.startswith("- ") for line in summary.splitlines())
            answer = summary

        right = (not found) if nothing else (wanted.casefold() in answer.casefold())
        passed += right
        print(f"  {'PASS' if right else 'FAIL'} {store:7} {question}")

        if not right:
            print(f"       wanted {'nothing' if nothing else repr(wanted)}; got: {answer[:160]!r}")

    if total:
        print(f"\n  {passed} of {total} of your questions answered as expected ({passed / total:.0%}).")


def _report_citations(report_search):
    """How well each saved report cites its sources: the faithfulness side of RAG."""
    try:
        from docx import Document
    except ImportError:
        return

    paths = report_search.report_files()

    if not paths:
        return

    print("\nYour reports' citations (reports written before citations began show none):\n")
    citation = re.compile(r"\[\d+(?:\s*[,\u2013-]\s*\d+)*\]")

    for path in paths:
        try:
            paragraphs = [paragraph.text for paragraph in Document(path).paragraphs]
        except Exception as error:
            print(f"  could not read {os.path.basename(path)}: {error}")
            continue

        body = []

        for text in paragraphs:
            if text.strip().casefold() in ("sources", "## sources", "# sources", "**sources**"):
                break

            body.append(text)

        sentences = [
            sentence for text in body
            if text.strip() and not text.lstrip().startswith(("#", "|", "---")) and len(text.split()) > 6
            for sentence in re.split(r"(?<=[.!?])\s+", text) if len(sentence.split()) > 4
        ]
        cited = sum(bool(citation.search(sentence)) for sentence in sentences)
        unverified = sum(sentence.count("(unverified") for sentence in sentences)
        label, day = report_search.describe_report(path)
        share = f"{cited / len(sentences):.0%}" if sentences else "-"
        print(f"  {label[:48]:48} {str(day):10}  {share:>4} of sentences cited, {unverified} unverified")


# ---- the result ----------------------------------------------------------------

def show(suites, baseline):
    worse = []

    print(f"\n{'':10}{'score':>16}{'baseline':>10}")

    for suite in suites:
        timing = suite.milliseconds()
        print(f"{suite.name}  ({timing} ms a search)" if timing is not None else suite.name)

        for metric, value in suite.scores.items():
            before = (baseline.get(suite.name) or {}).get(metric)
            mark = ""

            if value is not None and before is not None and value < before - TOLERANCE:
                mark = "  WORSE"
                worse.append(f"{suite.name} {metric}: {value} (baseline {before})")
            elif value is not None and before is not None and value > before + TOLERANCE:
                mark = "  better"

            shown_before = "-" if before is None else f"{before:.3f}"
            shown_value = "-" if value is None else f"{value:.3f}"
            print(f"  {metric:14}{shown_value:>10}{shown_before:>10}{mark}")

        for miss in suite.wrong:
            print(f"    miss: {miss}")

    return worse


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--update", action="store_true", help="save these scores as the new baseline")
    parser.add_argument("--mine", action="store_true", help="your own questions and reports, read-only")
    options = parser.parse_args()

    if options.mine:
        mine()
        return 0

    suites = benchmark()
    baseline = json.loads(BASELINE.read_text(encoding="utf-8")) if BASELINE.exists() else {}
    worse = show(suites, baseline)

    if options.update:
        BASELINE.write_text(
            json.dumps({suite.name: suite.scores for suite in suites}, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"\nBaseline saved to {BASELINE.relative_to(ROOT)}.")
        return 0

    if worse:
        print("\nWorse than the baseline:\n  " + "\n  ".join(worse))
        return 1

    print("\nNo score is worse than the baseline." if baseline else "\nNo baseline yet: run with --update to save one.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
